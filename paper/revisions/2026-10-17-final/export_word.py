"""Deterministic editable Word export of the final LaTeX source, not a new draft."""
from pathlib import Path
import copy, hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

O=Path(__file__).resolve().parent
tex=(O/'manuscript.tex').read_text(encoding='utf8')
labels=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',(O/'manuscript.aux').read_text()))
items=re.findall(r'\\bibitem\[(.*?)\]\{([^}]+)\}\s*(.*?)(?=\\bibitem|\\end\{thebibliography\})',(O/'manuscript.bbl').read_text(encoding='utf8'),re.S)
cites={key:re.sub(r'\s+',' ',label.split(')')[0]+')').replace('~',' ').replace('(', ' (') for label,key,_ in items}

def arg(s,i):
    assert s[i]=='{'
    depth=1;j=i+1
    while depth:
        if s[j]=='{' and s[j-1]!='\\':depth+=1
        if s[j]=='}' and s[j-1]!='\\':depth-=1
        j+=1
    return s[i+1:j-1],j

def plain(s,nested=False):
    s=s.replace(r'\allowbreak','').replace(r'\newline','; ')
    s=s.replace(r'{\={u}}','ū').replace(r'{\.{e}}','ė')
    for macro in ['citep','citet','ref','emph','texttt','textbf','url','doi','href']:
        while m:=re.search(r'\\'+macro+r'\{',s):
            a,end=arg(s,m.end()-1)
            if macro.startswith('cite'):
                vals=[cites[k.strip()] for k in a.split(',')]
                value='['+'; '.join(vals)+']' if macro=='citep' else '; '.join(vals)
            elif macro=='ref':value=labels[a]
            elif macro=='href':value,end=arg(s,end);value=plain(value,True)
            elif macro=='doi':value='doi: '+a
            else:value=plain(a,True)
            s=s[:m.start()]+value+s[end:]
    s=re.sub(r'\\label\{[^}]+\}','',s)
    s=s.replace(r'\newblock',' ').replace(r'\penalty0','')
    s=s.replace(r'\{','⦃').replace(r'\}','⦄')
    for a,b in [(r'\times','×'),(r'\%','%'),(r'\_','_'),(r'\&','&'),(r'\$','$'),(r'\,',' '),('---','—'),('--','–'),('~',' ')]:s=s.replace(a,b)
    s=s.replace('{','').replace('}','')
    if not nested:s=s.replace('⦃','{').replace('⦄','}')
    s=s.replace('``','“').replace("''",'”')
    return re.sub(r'\s+',' ',s).strip()

M='http://www.w3.org/1998/Math/MathML'
transform=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
def mathml(source):
    return copy.deepcopy(transform(etree.fromstring(('<math xmlns="'+M+'">'+source+'</math>').encode())).getroot())
def mi(s):return '<mi>'+s+'</mi>'
def sub(base,index):return '<msub>'+mi(base)+'<mi>'+index+'</mi></msub>'
inline={
 'c_{imXr}':sub('c','imXr'),
 '10{,}000\\times2\\times3\\times3=180{,}000':'<mrow><mn>10,000</mn><mo>×</mo><mn>2</mn><mo>×</mo><mn>3</mn><mo>×</mo><mn>3</mn><mo>=</mo><mn>180,000</mn></mrow>',
 'n=1':'<mi>n</mi><mo>=</mo><mn>1</mn>',
 '98+145-34=209':'<mn>98</mn><mo>+</mo><mn>145</mn><mo>−</mo><mn>34</mn><mo>=</mo><mn>209</mn>'}
for key in ('m','i','X','r','K'):inline[key]=mi(key)
def inline_math(raw):
    assert raw in inline,raw
    return mathml(inline[raw])
def eq(condition):
    right='Y' if condition==1 else 'X'
    lhs='<msup><msub><mi>L</mi><mi>im</mi></msub><mrow><mi>X</mi><mo>→</mo><mi>'+right+'</mi></mrow></msup><mo>=</mo>'
    power='<msup><mi>K</mi><mrow><mo>−</mo><mn>1</mn></mrow></msup>' if condition==1 else '<msup><mrow><mo>[</mo><mi>K</mi><mo>(</mo><mi>K</mi><mo>−</mo><mn>1</mn><mo>)</mo><mo>]</mo></mrow><mrow><mo>−</mo><mn>1</mn></mrow></msup>'
    lower='<mi>r</mi>' if condition==1 else '<mrow><mi>r</mi><mo>≠</mo><mi>s</mi></mrow>'
    factor='<mrow>'+sub('c','imXr')+'<mo>(</mo><mn>1</mn><mo>−</mo>'+sub('c','imYr' if condition==1 else 'imXs')+'<mo>)</mo></mrow>'
    summation='<munder><mo>∑</mo>'+lower+'</munder>'
    return mathml('<mrow>'+lhs+power+summation+factor+'</mrow>')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.9)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(5);st.font.bold=name.startswith('Heading');st.font.italic=False
    fonts=st.element.rPr.rFonts
    for key in list(fonts.attrib):
        if key.endswith('Theme'):del fonts.attrib[key]
    for border in list(st.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
doc.styles['Title'].font.size=Pt(19)
doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.line_spacing=1.03
doc.core_properties.title=plain(arg(tex,tex.index(r'\title{')+6)[0])
doc.core_properties.author='Chimdumebi Mitchell Nebolisa; Jessica Udry'
doc.core_properties.last_modified_by='AI-assisted manuscript revision'
for value,style in [(doc.core_properties.title,'Title'),(doc.core_properties.author,'Subtitle'),('October 17, 2026','Subtitle')]:doc.add_paragraph(value,style)
footer=sec.footer.paragraphs[0];footer.alignment=1
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
paragraphs={};tables=[];figures=[];references=[];quotes=[];equations=[]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def text_of(p):return ''.join(p.xpath('.//w:t/text() | .//m:t/text()',namespaces=ns))
def mixed(p,raw):
    pos=0;expected=''
    for m in re.finditer(r'(?<!\\)\$[^$]+\$|\\href\{',raw):
        if m.start()<pos:continue
        part=raw[pos:m.start()]
        value=(' ' if part[:1].isspace() else '')+plain(part)+(' ' if part[-1:].isspace() else '')
        if value:p.add_run(value);expected+=value
        if m[0].startswith('$'):
            om=inline_math(m[0][1:-1]);p._p.append(om);expected+=''.join(om.xpath('.//m:t/text()',namespaces=ns));pos=m.end()
        else:
            url,end=arg(raw,m.end()-1);label,end=arg(raw,end)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
            run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=plain(label);run.append(t);link.append(run);p._p.append(link);expected+=plain(label);pos=end
    part=raw[pos:];value=(' ' if part[:1].isspace() else '')+plain(part)
    if value:p.add_run(value);expected+=value
    return expected.strip()
def paragraph(raw,pid=None):
    p=doc.add_paragraph();expected=mixed(p,raw)
    if pid:paragraphs[pid]={'expected':expected,'source':raw}
    return p
def table(raw):
    cap=arg(raw,raw.index(r'\caption{')+8)[0];label=re.search(r'\\label\{([^}]+)\}',raw)[1]
    p=doc.add_paragraph('Table '+labels[label]+': '+plain(cap),'Caption');p.paragraph_format.keep_with_next=True
    grid=raw.split(r'\toprule',1)[1].split(r'\bottomrule',1)[0].replace(r'\midrule','')
    rows=[[plain(c) for c in row.strip().split('&')] for row in grid.split(r'\\') if row.strip()]
    n=len(rows[0]);widths={4:[.65,.65,1.25,4.15],3:[1.4,2.65,2.65],7:[2.3,.7,.7,.75,.75,.75,.75]}.get(n,[6.7/n]*n)
    if label=='tab:transitions':widths=[.45,.45,.85,.7,.7,.85,2.7]
    if label=='tab:variability':widths=[.65,.8,3.85,1.4]
    if label=='tab:examples':widths=[1.3,1.8,.8,2.8]
    if label=='tab:features':widths=[.8,.9,.9,.85,1.05,1.1,1.1]
    if label=='tab:alignment':widths=[1.8,.6,.6,.8,.8,1.1,1.0]
    t=doc.add_table(rows=0,cols=n);t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for ri,row in enumerate(rows):
        assert len(row)==n
        cells=t.add_row().cells
        for cell,w,value in zip(cells,widths,row):
            cell.width=Inches(w);cell.text=value
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9)
            props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'45');el.set(qn('w:type'),'dxa');margins.append(el)
            props.append(margins)
            if ri==0:
                for r in cell.paragraphs[0].runs:r.bold=True
        if ri==0:
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    tables.append({'label':label,'cells':rows})
def figure(raw):
    cap=arg(raw,raw.index(r'\caption{')+8)[0];label=re.search(r'\\label\{([^}]+)\}',raw)[1]
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(O/'figures/workflow.png'),width=Inches(6.7));shape._inline.docPr.set('descr',plain(cap))
    doc.add_paragraph('Figure '+labels[label]+': '+plain(cap),'Caption');figures.append({'label':label,'caption':plain(cap)})
body=tex.split(r'\maketitle',1)[1].split(r'\end{document}',1)[0]
pattern=re.compile(r'\\begin\{(figure|table|quote|align)\}.*?\\end\{\1\}|\\(?:sub)?section\{[^}]+\}(?:\\label\{[^}]+\})?|% P:[^\n]+\n[^\n]+|\\bibliography\{[^}]+\}|\\appendix|\\begin\{abstract\}',re.S)
section=0;subsection=0;appendix=False
for match in pattern.finditer(body):
    raw=match[0]
    if raw.startswith(r'\begin{table}'):table(raw)
    elif raw.startswith(r'\begin{figure}'):figure(raw)
    elif raw.startswith(r'\begin{quote}'):
        value=raw.split('\n',1)[1].split(r'\end{quote}')[0]
        value=value.replace(r'\\[4pt]', '\n\n');value=re.sub(r'\\\\(?:\[[^]]+\])?','\n',value).strip();value='\n'.join(plain(s) for s in value.splitlines())
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.25)
        r=p.add_run(value);r.font.name='Courier New';r.font.size=Pt(9);quotes.append(value)
    elif raw.startswith(r'\begin{align}'):
        assert r'c_{imYr}' in raw and r'c_{imXs}' in raw and r'K(K-1)' in raw
        for i in (1,2):
            p=doc.add_paragraph();p.alignment=1;p._p.append(eq(i));p.add_run('   ('+str(i)+')');equations.append(i)
    elif raw.startswith('% P:'):
        pid,content=raw.split('\n',1);paragraph(content,pid[4:])
    elif raw==r'\appendix':appendix=True;section=0
    elif raw.startswith(r'\begin{abstract}'):doc.add_paragraph('Abstract','Heading 1')
    elif raw.startswith(r'\bibliography'):
        doc.add_paragraph('References','Heading 1')
        for _,key,content in items:
            p=paragraph(content);p.paragraph_format.left_indent=Inches(.15);p.paragraph_format.first_line_indent=Inches(-.15);references.append({'id':key,'text':plain(content)})
    elif raw.startswith(r'\subsection'):
        subsection+=1;doc.add_paragraph(str(section)+'.'+str(subsection)+' '+plain(arg(raw,raw.index('{'))[0]),'Heading 2')
    else:
        section+=1;subsection=0;prefix=chr(64+section) if appendix else str(section)
        heading=doc.add_paragraph(prefix+' '+plain(arg(raw,raw.index('{'))[0]),'Heading 1')
        if appendix and section==1:heading.paragraph_format.page_break_before=True
doc.save(O/'manuscript.docx')
with zipfile.ZipFile(O/'manuscript.docx') as z:
    root=etree.fromstring(z.read('word/document.xml'))
    actual=[text_of(p) for p in root.xpath('//w:body/w:p',namespaces=ns)]
    norm=lambda s:re.sub(r'\s+',' ',s).strip()
    for pid,r in paragraphs.items():r['equal']=norm(r['expected']) in [norm(a) for a in actual]
    alltext='\n'.join(actual)
    checks={'all_source_paragraphs':len(paragraphs)==48 and all(r['equal'] for r in paragraphs.values()),
      'nine_native_tables':len(doc.tables)==9,'one_figure':len(doc.inline_shapes)==1,
      'nine_references':len(references)==9 and all(r['text'] in actual for r in references),
      'literal_evidence_placeholder':any('Evidence: {gold_evidence}' in a for a in actual),
      'two_exact_message_quotes':len(quotes)==2 and '{claim_text}' in quotes[0] and 'based only on the evidence provided.' in quotes[1],
      'two_native_display_equations':equations==[1,2] and len(root.xpath('//m:oMath',namespaces=ns))==11,
      'all_native_table_values':all([[c.text for c in row.cells] for row in t.rows]==r['cells'] for t,r in zip(doc.tables,tables)),
      'no_unconverted_commands':not re.search(r'\\(?:cite|ref|emph|texttt|times|begin|end|[A-Za-z]+)',alltext),
      'valid_zip':z.testzip() is None}
report={'status':'PASS' if all(checks.values()) else 'FAIL','checks':checks,'master_sha256':hashlib.sha256((O/'manuscript.tex').read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256((O/'manuscript.docx').read_bytes()).hexdigest(),'paragraphs':paragraphs,'tables':tables,'figures':figures,'references':references,'quotes':quotes,'equations':equations,'visual_status':'NOT_RENDERED','equation_method':'Native OMML transformed from explicitly source-matched MathML with installed Office MML2OMML.XSL; visual QA required.'}
(O/'reports/word-equivalence.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf8')
print(report['status'],checks)
raise SystemExit(0 if report['status']=='PASS' else 1)
